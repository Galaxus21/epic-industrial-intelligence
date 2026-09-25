"""
EPIC — The text of an uploaded file, for entity extraction and search.

One reader per file type. A PDF keeps its page boundaries, headings and tables as markdown, so it may run to
PDF_TEXT_CHARS; every other type is cut at TEXT_CHARS. A file that cannot be read yields "", except a presentation,
which yields a bracketed note the pipeline does not index.
"""
from __future__ import annotations

import json
import logging
import os
from typing import Callable

logger = logging.getLogger(__name__)

TEXT_CHARS = 8000
# Page markers, table pipes and headings take room, so a PDF may use more before it is cut.
PDF_TEXT_CHARS = 32000
BYTES_PER_KB = 1024


def extractText(path: str, ext: str) -> str:
    """The text of the file at `path`, read by its extension `ext`; "" when it cannot be read."""
    reader = TEXT_READERS.get(ext)
    if reader is None:
        return ""
    try:
        return reader(path)
    except Exception as exc:
        logger.warning("Text extraction failed for %s: %s", path, exc)
        return ""


def _pdfText(path: str) -> str:
    import fitz
    import pymupdf4llm

    try:
        chunks = pymupdf4llm.to_markdown(path, page_chunks=True)
        pages = [(chunk.get("metadata", {}).get("page_number") or index + 1, (chunk.get("text") or "").strip())
                 for index, chunk in enumerate(chunks)]
        text = "\n\n".join(f"--- Page {number} ---\n{pageText}" for number, pageText in pages if pageText)
    except Exception as exc:
        logger.warning("pymupdf4llm extraction failed for %s: %s, falling back to fitz", path, exc)
        document = fitz.open(path)
        text = "\n\n".join(f"--- Page {index + 1} ---\n{page.get_text()}" for index, page in enumerate(document))
    return text[:PDF_TEXT_CHARS]


def _docxText(path: str) -> str:
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)[:TEXT_CHARS]


def _plainText(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace") as file:
        return file.read(TEXT_CHARS)


def _xlsxText(path: str) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    rows = ["\t".join(str(cell) for cell in row if cell is not None)
            for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
    return "\n".join(rows)[:TEXT_CHARS]


def _jsonText(path: str) -> str:
    raw = _plainText(path)
    try:
        return json.dumps(json.loads(raw), indent=2)[:TEXT_CHARS]
    except ValueError:
        return raw


def _pptxText(path: str) -> str:
    try:
        from pptx import Presentation

        lines: list[str] = []
        for number, slide in enumerate(Presentation(path).slides, 1):
            lines.append(f"--- Slide {number} ---")
            lines.extend(shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
        return "\n".join(lines)[:TEXT_CHARS]
    except Exception as exc:
        logger.warning("PPTX extraction failed for %s: %s", path, exc, exc_info=True)
        return f"[PPTX PRESENTATION — {os.path.getsize(path) // BYTES_PER_KB} KB — text extraction failed]"


TEXT_READERS: dict[str, Callable[[str], str]] = {
    ".pdf": _pdfText,
    ".docx": _docxText,
    ".txt": _plainText,
    ".csv": _plainText,
    ".xlsx": _xlsxText,
    ".json": _jsonText,
    ".pptx": _pptxText,
}
